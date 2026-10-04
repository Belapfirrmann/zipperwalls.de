"""Liest einen Newsletter-Plan (xlsx) ein und schreibt src/newsletterplan.js.

Aufruf:  pip install openpyxl && python3 scripts/import_newsletter.py plans/Newsletter-Plan_Q4_2026.xlsx
Erwartet Blatt "Newsletter-Plan" mit Kopfzeile: Nr | KW | Datum | Tag | Uhrzeit | Typ | Bezug Social Media | Rubrik |
Thema | Betreff | Betreff Alternative (A/B) | Vorschautext | Text | Button-Text | Button-Link | Zusatzblock Titel |
Zusatzblock Text | Zusatzblock Link | Bildidee | Hinweis / offen | Status | Öffnungsrate % | Klickrate % | Abmeldungen
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

COLS = {
    'Nr': 'nr', 'KW': 'kw', 'Datum': 'date', 'Tag': 'day', 'Uhrzeit': 'time', 'Typ': 'type',
    'Bezug Social Media': 'social', 'Rubrik': 'rubric', 'Thema': 'topic', 'Betreff': 'subject',
    'Betreff Alternative (A/B)': 'subject_alt', 'Vorschautext': 'preview', 'Text': 'text',
    'Button-Text': 'button_text', 'Button-Link': 'button_link', 'Zusatzblock Titel': 'extra_title',
    'Zusatzblock Text': 'extra_text', 'Zusatzblock Link': 'extra_link', 'Bildidee': 'image_idea',
    'Hinweis / offen': 'note', 'Status': 'status', 'Öffnungsrate %': 'open_rate', 'Klickrate %': 'click_rate',
    'Abmeldungen': 'unsubscribes',
}


def clean(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M')
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        if re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', v):
            d, m, y = v.split('.')
            return f'{y}-{m}-{d}'
        return v or None
    return v


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Newsletter-Plan']
    rows = list(ws.iter_rows(values_only=True))
    head_idx = next(i for i, r in enumerate(rows) if r and r[0] == 'Nr')
    header = [COLS.get(clean(h)) for h in rows[head_idx]]
    entries = []
    for r in rows[head_idx + 1:]:
        if r[0] is None:
            continue
        e = {k: clean(v) for k, v in zip(header, r) if k}
        e['nr'] = int(e['nr'])
        e['kw'] = int(e['kw']) if e.get('kw') is not None else None
        e['time'] = e.get('time') or '09:00'
        e['status'] = e.get('status') or 'Entwurf'
        entries.append(e)

    # Legende: Titelzeile, Quartalszeile, Arbeitsablauf und Quellen
    legend = [clean(r[0]) for r in wb['Legende'].iter_rows(values_only=True)] if 'Legende' in wb.sheetnames else []
    legend = [x for x in legend if x]
    quarter = next((x for x in legend if x.startswith('Quartal:')), None)
    sources, steps, in_sources = [], [], False
    for x in legend:
        if x == 'Quellen':
            in_sources = True
            continue
        if in_sources:
            sources.append(x)
        elif re.match(r'^\d\. ', x):
            steps.append(x)

    m = re.search(r'Q(\d)[_ ](\d{4})', Path(path).name)
    title = f'Newsletter-Plan Q{m.group(1)} {m.group(2)}' if m else 'Newsletter-Plan'
    types = {}
    for e in entries:
        types[e['type']] = types.get(e['type'], 0) + 1

    plan = {
        'title': title,
        'intro': quarter,
        'source': Path(path).name,
        'stand': datetime.date.today().isoformat(),
        'steps': steps,
        'sources': sources,
        'types': [{'name': k, 'count': v} for k, v in sorted(types.items(), key=lambda x: -x[1])],
        'entries': entries,
    }
    out = Path(__file__).resolve().parent.parent / 'src' / 'newsletterplan.js'
    out.write_text(
        '// Automatisch erzeugt von scripts/import_newsletter.py. Nicht von Hand bearbeiten.\n'
        'export const NEWSLETTERPLAN = ' + json.dumps(plan, ensure_ascii=False, indent=2) + ';\n',
        encoding='utf-8',
    )
    print(f'{len(entries)} Newsletter nach {out} geschrieben')


if __name__ == '__main__':
    main(sys.argv[1])
