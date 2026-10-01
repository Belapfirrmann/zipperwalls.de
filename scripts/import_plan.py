"""Liest einen Redaktionsplan (xlsx) ein und schreibt src/redaktionsplan.js.

Aufruf:  pip install openpyxl && python3 scripts/import_plan.py plans/Redaktionsplan_Social_Media_Q4_2026.xlsx
Erwartet Blatt "Redaktionsplan" mit Kopfzeile: Nr | KW | Datum | Tag | Pillar | Thema | Format |
Bildtext (Headline) | Instagram | Facebook | LinkedIn | Hashtags | Grafik-Dateien | Hinweis / offen | Status
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

COLS = {
    'Nr': 'nr', 'KW': 'kw', 'Datum': 'date', 'Tag': 'day', 'Pillar': 'pillar', 'Thema': 'topic',
    'Format': 'format', 'Bildtext (Headline)': 'headline', 'Instagram': 'instagram', 'Facebook': 'facebook',
    'LinkedIn': 'linkedin', 'Hashtags': 'hashtags', 'Grafik-Dateien': 'files', 'Hinweis / offen': 'note', 'Status': 'status',
}


def clean(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip() or None
    return v


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Redaktionsplan']
    rows = list(ws.iter_rows(values_only=True))
    title = clean(rows[0][0]) or 'Redaktionsplan'
    intro = clean(rows[1][0])
    head_idx = next(i for i, r in enumerate(rows) if r and r[0] == 'Nr')
    header = [COLS.get(clean(h)) for h in rows[head_idx]]
    entries = []
    for r in rows[head_idx + 1:]:
        if r[0] is None:
            continue
        e = {k: clean(v) for k, v in zip(header, r) if k}
        e['hashtags'] = re.findall(r'#[\wäöüÄÖÜß]+', e.get('hashtags') or '')
        e['files'] = [f.strip() for f in (e.get('files') or '').splitlines() if f.strip()]
        entries.append(e)

    legend = {}
    if 'Legende' in wb.sheetnames:
        for r in wb['Legende'].iter_rows(min_row=2, values_only=True):
            if r[0] and r[1]:
                legend[clean(r[0])] = clean(r[1])

    pillars = {}
    for e in entries:
        pillars[e['pillar']] = pillars.get(e['pillar'], 0) + 1

    plan = {
        'title': title,
        'intro': intro,
        'source': Path(path).name,
        'stand': datetime.date.today().isoformat(),
        'rhythm': {'postsPerWeek': 2, 'days': ['Dienstag', 'Donnerstag'], 'time': '09:00', 'window': legend.get('Posting-Zeit')},
        'formats': legend.get('Formate'),
        'legend': legend,
        'pillars': [{'name': k, 'count': v} for k, v in sorted(pillars.items(), key=lambda x: -x[1])],
        'entries': entries,
    }
    out = Path(__file__).resolve().parent.parent / 'src' / 'redaktionsplan.js'
    out.write_text(
        '// Automatisch erzeugt von scripts/import_plan.py. Nicht von Hand bearbeiten.\n'
        'export const REDAKTIONSPLAN = ' + json.dumps(plan, ensure_ascii=False, indent=2) + ';\n',
        encoding='utf-8',
    )
    print(f'{len(entries)} Einträge nach {out} geschrieben')


if __name__ == '__main__':
    main(sys.argv[1])
